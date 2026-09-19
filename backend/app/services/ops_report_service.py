"""N7 运维报告智能生成：聚合既有运维数据 → LLM 结构化报告 → 落库/导出。

数据源全部来自既有 MySQL 表（不新采数据）：
- RealtimeEvent 告警事件 / ProactiveOpsRun 主动运维处置
- Ticket 两票 / OpsSnapshot 交接班快照 / TelemetryPoint 遥测
生成链路：collect_data → LLM 按 JSON 分节生成 → 失败降级为纯模板拼接（degraded 不 crash）。
"""
import asyncio
import json
import uuid
from datetime import datetime, timedelta

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.core.obs import degraded
from app.core.response import BizError
from app.models.ops_report import OpsReport
from app.models.ops_workbench import OpsSnapshot, TelemetryPoint
from app.models.realtime_event import ProactiveOpsRun, RealtimeEvent
from app.models.ticket import Ticket

REPORT_TYPES = {
    "shift": {"label": "交接班报告", "days": 1,
              "sections": ["本班概览", "告警与事件", "两票执行情况", "遥测异动", "交接事项与建议"]},
    "weekly": {"label": "运维周报", "days": 7,
               "sections": ["本周概览", "告警趋势", "两票统计", "设备健康与风险", "下周建议"]},
    "fault": {"label": "故障分析报告", "days": 7,
              "sections": ["故障概况", "事件时间线", "处置过程", "原因初判", "后续措施建议"]},
    "device": {"label": "设备健康报告", "days": 7,
               "sections": ["设备概览", "遥测分析", "关联告警", "关联两票", "健康结论"]},
}

_MAX_ROWS = 500  # 单数据源拉取上限（防大范围全表扫描）


def _row(r: OpsReport, *, with_content: bool = False) -> dict:
    d = {"id": r.id, "title": r.title, "reportType": r.report_type, "status": r.status,
         "params": r.params or {}, "summary": r.summary, "error": r.error,
         "llmMeta": r.llm_meta or {}, "createdBy": r.created_by,
         "createdAt": r.created_at.isoformat() if r.created_at else None,
         "updatedAt": r.updated_at.isoformat() if r.updated_at else None}
    if with_content:
        d["dataSnapshot"] = r.data_snapshot or {}
        d["sections"] = r.sections or []
        d["contentMd"] = r.content_md
    return d


# ---------- 数据聚合（纯 SQL，租户隔离） ----------

async def _collect_alarms(db: AsyncSession, tenant: str, start: datetime) -> dict:
    base = select(RealtimeEvent).where(
        RealtimeEvent.tenant_id == tenant, RealtimeEvent.occurred_at >= start)
    rows = (await db.execute(base.order_by(RealtimeEvent.occurred_at.desc())
                             .limit(_MAX_ROWS))).scalars().all()
    by_sev: dict[str, int] = {}
    by_device: dict[str, int] = {}
    for r in rows:
        by_sev[r.severity] = by_sev.get(r.severity, 0) + 1
        dev = r.canonical_device_name or r.source_device_id or "未知设备"
        by_device[dev] = by_device.get(dev, 0) + 1
    return {
        "total": len(rows),
        "bySeverity": by_sev,
        "topDevices": sorted(by_device.items(), key=lambda kv: -kv[1])[:10],
        "recent": [{"time": r.occurred_at.isoformat() if r.occurred_at else "",
                    "severity": r.severity, "title": r.title,
                    "device": r.canonical_device_name or r.source_device_id}
                   for r in rows[:20]],
    }


async def _collect_tickets(db: AsyncSession, tenant: str, start: datetime) -> dict:
    rows = (await db.execute(
        select(Ticket).where(Ticket.tenant_id == tenant, Ticket.is_deleted == 0,
                             Ticket.created_at >= start).limit(_MAX_ROWS))).scalars().all()
    by_status: dict[str, int] = {}
    by_type: dict[str, int] = {}
    for t in rows:
        st = t.status.value if hasattr(t.status, "value") else str(t.status)
        by_status[st] = by_status.get(st, 0) + 1
        tp = t.ticket_type.value if hasattr(t.ticket_type, "value") else str(t.ticket_type)
        by_type[tp] = by_type.get(tp, 0) + 1
    return {"total": len(rows), "byStatus": by_status, "byType": by_type,
            "recent": [{"id": t.id,
                        "type": t.ticket_type.value if hasattr(t.ticket_type, "value") else str(t.ticket_type),
                        "status": t.status.value if hasattr(t.status, "value") else str(t.status),
                        "task": (t.title or "")[:80]}
                       for t in rows[:15]]}


async def _collect_proactive(db: AsyncSession, tenant: str, start: datetime) -> dict:
    rows = (await db.execute(
        select(ProactiveOpsRun).where(ProactiveOpsRun.tenant_id == tenant,
                                      ProactiveOpsRun.created_at >= start)
        .limit(_MAX_ROWS))).scalars().all()
    by_status: dict[str, int] = {}
    for r in rows:
        by_status[r.status] = by_status.get(r.status, 0) + 1
    return {"total": len(rows), "byStatus": by_status}


async def _collect_handovers(db: AsyncSession, tenant: str, start: datetime) -> list[dict]:
    rows = (await db.execute(
        select(OpsSnapshot).where(OpsSnapshot.tenant_id == tenant,
                                  OpsSnapshot.created_at >= start)
        .order_by(OpsSnapshot.created_at.desc()).limit(20))).scalars().all()
    out = []
    for s in rows:
        payload: dict = {}
        try:
            payload = json.loads(s.payload_json) if s.payload_json else {}
        except Exception as e:  # noqa: BLE001 — 快照损坏跳过该条，不阻断报告
            degraded("ops_report_snapshot_parse", e, s.id)
        out.append({"kind": s.kind, "title": s.title, "creator": s.creator,
                    "createdAt": s.created_at.isoformat() if s.created_at else None,
                    "payload": payload})
    return out


async def _collect_telemetry(db: AsyncSession, tenant: str, start: datetime,
                             device_id: str = "") -> dict:
    conds = [TelemetryPoint.tenant_id == tenant, TelemetryPoint.observed_at >= start]
    if device_id:
        conds.append(TelemetryPoint.device_id == device_id)
    # 按设备+指标聚合统计值 + 最近观测（取最新 2000 点在内存归并，避免窗口函数跨 sqlite/MySQL 差异）
    rows = (await db.execute(
        select(TelemetryPoint.device_id, TelemetryPoint.metric, TelemetryPoint.unit,
               TelemetryPoint.value, TelemetryPoint.observed_at)
        .where(*conds).order_by(TelemetryPoint.observed_at.desc()).limit(2000))).all()
    stats: dict[tuple[str, str], dict] = {}
    for dev, metric, unit, value, observed in rows:
        key = (dev, metric)
        st = stats.setdefault(key, {"min": value, "max": value, "sum": 0.0, "n": 0,
                                    "unit": unit, "latest": value,
                                    "latestAt": observed.isoformat() if observed else ""})
        st["min"] = min(st["min"], value)
        st["max"] = max(st["max"], value)
        st["sum"] += value
        st["n"] += 1
    metrics = [{"device": k[0], "metric": k[1], **{kk: vv for kk, vv in v.items()}}
               for k, v in sorted(stats.items())]
    for m in metrics:
        m["avg"] = round(m.pop("sum") / m["n"], 3) if m["n"] else 0
    return {"metrics": metrics[:40]}


async def collect_data(db: AsyncSession, tenant: str, report_type: str, days: int,
                       device_id: str = "") -> dict:
    """聚合报告数据事实。任一数据源异常降级为空（报告继续，标注 degraded）。"""
    start = datetime.now() - timedelta(days=max(1, days))
    data: dict = {"windowStart": start.isoformat(), "windowDays": days,
                  "collectedAt": datetime.now().isoformat(), "sources": {}}
    collectors = {
        "alarms": lambda: _collect_alarms(db, tenant, start),
        "tickets": lambda: _collect_tickets(db, tenant, start),
        "proactive": lambda: _collect_proactive(db, tenant, start),
        "handovers": lambda: _collect_handovers(db, tenant, start),
        "telemetry": lambda: _collect_telemetry(db, tenant, start, device_id),
    }
    for name, fn in collectors.items():
        try:
            data["sources"][name] = await fn()
        except Exception as e:  # noqa: BLE001 — 单源失败不影响整份报告
            degraded(f"ops_report_collect_{name}", e)
            data["sources"][name] = {"error": str(e)}
    return data


# ---------- LLM 生成 + 模板降级 ----------

_GEN_PROMPT = """你是电网运维报告专家。基于下面的运维数据事实，生成一份{label}。
要求：
- 只依据数据事实，不编造；数据缺失的条目如实写"本期无数据"
- 每节 2-6 句，运维术语规范，结论可执行
- 输出严格 JSON：{{"sections":[{{"title":"...","content":"..."}}],"summary":"一句话摘要"}}
- sections 的 title 必须依次为：{titles}
- summary 不超过 60 字

数据事实（JSON）：
{data}"""


async def _generate_sections(report_type: str, data: dict,
                             model_type: str | None = None) -> dict:
    """LLM 生成分节；失败抛给调用方降级。"""
    from app.providers.factory import get_llm_provider

    spec = REPORT_TYPES[report_type]
    prompt = _GEN_PROMPT.format(label=spec["label"], titles="、".join(spec["sections"]),
                                data=json.dumps(data["sources"], ensure_ascii=False)[:12000])
    provider = get_llm_provider(model_type)
    raw = await provider.chat([{"role": "user", "content": prompt}],
                              temperature=0.3, max_tokens=1500)
    text = (raw or "").strip()
    # 容错：模型可能裹 ```json 围栏
    if text.startswith("```"):
        text = text.strip("`").lstrip("json").strip()
    end = text.rfind("}")
    parsed = json.loads(text[: end + 1] if end >= 0 else text)
    sections = [{"title": str(s.get("title", "")), "content": str(s.get("content", ""))}
                for s in parsed.get("sections", []) if s.get("title")]
    if not sections:
        raise ValueError("LLM 未返回有效 sections")
    return {"sections": sections, "summary": str(parsed.get("summary", ""))[:120],
            "llmMeta": {"model": model_type or "default", "promptChars": len(prompt)}}


def _template_sections(report_type: str, data: dict) -> dict:
    """LLM 不可用时的确定性模板拼接（数据事实直陈，保证报告可交付）。"""
    src = data["sources"]
    alarms = src.get("alarms", {})
    tickets = src.get("tickets", {})
    proactive = src.get("proactive", {})
    handovers = src.get("handovers", [])
    telemetry = src.get("telemetry", {})
    parts: list[dict] = []

    def _sev(alarms_dict: dict) -> str:
        sev = alarms_dict.get("bySeverity", {})
        return "，".join(f"{k} {v} 条" for k, v in sorted(sev.items())) or "本期无告警"

    if report_type in ("shift", "weekly"):
        parts.append({"title": REPORT_TYPES[report_type]["sections"][0],
                      "content": f"统计窗口近 {data['windowDays']} 天。"
                                 f"告警共 {alarms.get('total', 0)} 条（{_sev(alarms)}）；"
                                 f"新增两票 {tickets.get('total', 0)} 张；"
                                 f"主动运维处置 {proactive.get('total', 0)} 次。"})
    if report_type == "fault":
        recent = alarms.get("recent", [])
        tl = "；".join(f"{e.get('time', '')[-8:]} {e.get('device', '')} {e.get('title', '')}"
                       for e in recent[:8]) or "本期无告警记录"
        parts.append({"title": "故障概况",
                      "content": f"窗口内告警 {alarms.get('total', 0)} 条，"
                                 f"最集中设备：{alarms.get('topDevices') or '无'}。"})
        parts.append({"title": "事件时间线", "content": tl})
        parts.append({"title": "处置过程",
                      "content": f"主动运维共 {proactive.get('total', 0)} 次"
                                 f"（{proactive.get('byStatus', {})}）。"})
        parts.append({"title": "原因初判", "content": "（LLM 不可用，原因分析降级——"
                                                    "请结合时间线人工研判）"})
        parts.append({"title": "后续措施建议", "content": "复核高频告警设备遥测趋势，必要时安排特巡。"})
    mets = telemetry.get("metrics", [])
    if mets:
        m_txt = "；".join(f"{m['device']}/{m['metric']} 均值 {m['avg']}{m['unit']} "
                          f"(min {m['min']}~max {m['max']})" for m in mets[:8])
        parts.append({"title": "遥测异动" if report_type == "shift" else "遥测分析",
                      "content": m_txt})
    if handovers:
        parts.append({"title": "交接班记录",
                      "content": "；".join(f"{h.get('createdAt', '')[:16]} {h.get('title', '')}"
                                           for h in handovers[:5])})
    parts.append({"title": "风险与建议",
                  "content": "对告警集中设备加密监视；两票积压状态及时催办；本报告为模板降级版，"
                             "建议恢复 LLM 后重新生成。"})
    return {"sections": parts, "summary": f"模板降级版：告警 {alarms.get('total', 0)} 条，"
                                          f"两票 {tickets.get('total', 0)} 张",
            "llmMeta": {"fallback": True}}


def _assemble(title: str, report_type: str, data: dict, sections: list[dict]) -> str:
    lines = [f"# {title}", "",
             f"> 类型：{REPORT_TYPES[report_type]['label']} · "
             f"统计窗口：近 {data['windowDays']} 天（起 {data['windowStart'][:19]}） · "
             f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ""]
    for s in sections:
        lines += [f"## {s['title']}", "", s["content"], ""]
    lines.append("---")
    lines.append("⚠ 本报告由系统聚合数据自动生成，仅供运维辅助参考；现场操作前必须核对调度指令与安规。")
    return "\n".join(lines)


# ---------- 对外编排 ----------

async def create_report(db: AsyncSession, tenant_id: str, username: str,
                        report_type: str, days: int | None = None,
                        device_id: str = "", model_type: str | None = None,
                        title: str = "") -> dict:
    if report_type not in REPORT_TYPES:
        raise BizError(f"未知报告类型：{report_type}", 400)
    spec = REPORT_TYPES[report_type]
    d = min(max(days or spec["days"], 1), int(getattr(settings, "OPS_REPORT_MAX_DAYS", 31)))
    params = {"days": d, "deviceId": device_id}
    now = datetime.now().strftime("%Y%m%d_%H%M")
    rpt = OpsReport(
        id=uuid.uuid4().hex,
        title=(title or f"{spec['label']}·{now}").strip()[:120],
        report_type=report_type, status="generating", params=params,
        tenant=tenant_id, created_by=username)
    db.add(rpt)
    await db.commit()
    await db.refresh(rpt)  # MySQL 无 RETURNING，server_default 列需 refresh 才可读（否则懒 IO 崩）
    try:
        asyncio.create_task(_generate(rpt.id, report_type, params, tenant_id,
                                      model_type, username))
    except Exception as e:  # create_task 拒绝（事件循环关闭等）→ 直接同步生成
        degraded("ops_report_schedule", e)
        await _generate(rpt.id, report_type, params, tenant_id, model_type, username)
    return _row(rpt)


async def _generate(report_id: str, report_type: str, params: dict,
                    tenant_id: str, model_type: str | None, username: str,
                    db: AsyncSession | None = None) -> None:
    """后台生成：默认独立 session（与请求生命周期解耦），测试可注入；失败落 failed 状态。"""
    from app.db.session import AsyncSessionLocal

    async def _run(db: AsyncSession) -> None:
        rpt = (await db.execute(select(OpsReport).where(
            OpsReport.id == report_id, OpsReport.tenant == tenant_id))).scalar_one_or_none()
        if rpt is None:
            return
        try:
            data = await collect_data(db, tenant_id, report_type, params["days"],
                                      params.get("deviceId", ""))
            gen = None
            try:
                gen = await _generate_sections(report_type, data, model_type)
            except Exception as e:  # noqa: BLE001 — LLM 失败降级模板，报告仍可交付
                degraded("ops_report_llm", e)
                gen = _template_sections(report_type, data)
            rpt.data_snapshot = data
            rpt.sections = gen["sections"]
            rpt.summary = gen["summary"]
            rpt.llm_meta = gen["llmMeta"]
            rpt.content_md = _assemble(rpt.title, report_type, data, gen["sections"])
            rpt.status = "done"
        except Exception as e:  # noqa: BLE001 — 聚合阶段异常 → failed 可见，不静默
            degraded("ops_report_generate", e)
            rpt.status = "failed"
            rpt.error = str(e)
        await db.commit()

    if db is not None:
        await _run(db)
        return
    async with AsyncSessionLocal() as session:
        await _run(session)


async def regenerate(db: AsyncSession, report_id: str, tenant_id: str,
                     model_type: str | None = None) -> dict:
    rpt = await get_report(db, report_id, tenant_id)
    if rpt.status == "generating":
        raise BizError("报告正在生成中，请稍候", 400)
    rpt.status = "generating"
    rpt.error = ""
    await db.commit()
    await db.refresh(rpt)  # onupdate(updated_at) 后属性过期，async session 禁属性级懒 IO
    try:
        asyncio.create_task(_generate(rpt.id, rpt.report_type, rpt.params,
                                      tenant_id, model_type, rpt.created_by))
    except Exception as e:  # noqa: BLE001
        degraded("ops_report_schedule", e)
        await _generate(rpt.id, rpt.report_type, rpt.params, tenant_id,
                        model_type, rpt.created_by)
    return _row(rpt)


async def get_report(db: AsyncSession, report_id: str, tenant_id: str) -> OpsReport:
    rpt = (await db.execute(select(OpsReport).where(
        OpsReport.id == report_id, OpsReport.tenant == tenant_id))).scalar_one_or_none()
    if rpt is None:
        raise BizError("报告不存在", 404)
    return rpt


async def get_report_detail(db: AsyncSession, report_id: str, tenant_id: str) -> dict:
    return _row(await get_report(db, report_id, tenant_id), with_content=True)


async def list_reports(db: AsyncSession, tenant_id: str, report_type: str = "",
                       status: str = "", page: int = 1, size: int = 20) -> tuple[int, list[dict]]:
    conds = [OpsReport.tenant == tenant_id]
    if report_type:
        conds.append(OpsReport.report_type == report_type)
    if status:
        conds.append(OpsReport.status == status)
    total = (await db.execute(select(func.count(OpsReport.id)).where(*conds))).scalar() or 0
    rows = (await db.execute(
        select(OpsReport).where(*conds)
        .order_by(OpsReport.created_at.desc())
        .offset((page - 1) * size).limit(size))).scalars().all()
    return total, [_row(r) for r in rows]


async def delete_report(db: AsyncSession, report_id: str, tenant_id: str) -> dict:
    rpt = await get_report(db, report_id, tenant_id)
    await db.delete(rpt)
    await db.commit()
    return {"id": report_id}


def build_report_docx(rpt_dict: dict) -> bytes:
    """报告 → Word（复用 python-docx，现场打印归档）。"""
    import io

    from docx import Document

    doc = Document()
    doc.add_heading(rpt_dict.get("title", "运维报告"), level=0)
    params = rpt_dict.get("params", {})
    doc.add_paragraph(
        f"类型：{REPORT_TYPES.get(rpt_dict.get('reportType', ''), {}).get('label', '')} · "
        f"统计窗口：近 {params.get('days', '-')} 天 · 生成：{rpt_dict.get('createdBy', '')} "
        f"{(rpt_dict.get('createdAt') or '')[:19]}")
    if rpt_dict.get("summary"):
        doc.add_paragraph(f"摘要：{rpt_dict['summary']}")
    for s in rpt_dict.get("sections", []):
        doc.add_heading(s.get("title", ""), level=1)
        for line in (s.get("content", "") or "").split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
    doc.add_paragraph()
    doc.add_paragraph("⚠ 本报告由系统聚合数据自动生成，仅供运维辅助参考；"
                      "现场操作前必须核对调度指令与安规。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def type_meta() -> list[dict]:
    """前端下拉：报告类型元数据。"""
    return [{"type": k, "label": v["label"], "defaultDays": v["days"],
             "sections": v["sections"]} for k, v in REPORT_TYPES.items()]
